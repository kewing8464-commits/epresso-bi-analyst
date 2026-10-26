"""$1,000 test: copy the Sales workbook and add exactly $1,000 to ONE revenue cell.

Usage:
    python scripts/make_modified_data.py <Support_EX_B_Sales.xlsx> <Sales ID> [out.xlsx]

Only the Revenue cell in the row with that Sales ID changes; every other sheet, row and
value is copied untouched. Upload the new file in place of the original to test the app.
"""
import sys

from openpyxl import load_workbook


def main(src, sales_id, out=None):
    out = out or src.replace('.xlsx', '_plus1000.xlsx')
    wb = load_workbook(src)
    ws = wb['Sales']
    header = [c.value for c in ws[1]]
    id_col, rev_col = header.index('Sales ID') + 1, header.index('Revenue') + 1
    for row in range(2, ws.max_row + 1):
        if ws.cell(row, id_col).value == sales_id:
            details = {h: ws.cell(row, i + 1).value for i, h in enumerate(header)}
            print(f'Row {row}: {details}')
            cell = ws.cell(row, rev_col)
            original = cell.value
            cell.value = original + 1000
            wb.save(out)
            print(f'Original revenue: {original:,.2f}')
            print(f'Modified revenue: {cell.value:,.2f}')
            print('Expected effect : total revenue +1,000.00 and only the filters that '
                  'include this row change')
            print(f'Saved: {out}')
            return
    sys.exit(f'Sales ID {sales_id} not found')


if __name__ == '__main__':
    if len(sys.argv) < 3:
        sys.exit(__doc__)
    main(*sys.argv[1:])
