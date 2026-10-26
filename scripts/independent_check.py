"""Independent check of the app's key numbers.

Reads the raw files with openpyxl and the csv module only (no pandas, and none of the
app's code), then sums revenue with plain Python loops. If these totals match the
app's Validation tab, the app's loading, typing and joins did not change revenue.

Usage:
    python scripts/independent_check.py <Support_EX_B_Sales.xlsx> <Support_EX_B_Stores.csv>
"""
import csv
import sys
from collections import defaultdict

from openpyxl import load_workbook


def main(sales_path, stores_path):
    with open(stores_path, newline='', encoding='utf-8-sig') as fh:
        store_type = {r['ID']: r['StoreType'] for r in csv.DictReader(fh, delimiter='\t')}

    ws = load_workbook(sales_path, read_only=True)['Sales']
    rows = ws.iter_rows(values_only=True)
    header = list(next(rows))
    i_date, i_store, i_rev = header.index('Date'), header.index('Store ID'), header.index('Revenue')

    total, n = 0.0, 0
    by_channel, by_year, by_month = defaultdict(float), defaultdict(float), defaultdict(float)
    for r in rows:
        rev = float(r[i_rev])
        date = str(r[i_date])[:10]          # 'YYYY-MM-DD'
        total += rev
        n += 1
        by_channel[store_type.get(r[i_store], 'Unknown store')] += rev
        by_year[date[:4]] += rev
        by_month[date[5:7]] += rev

    print(f'Rows read: {n:,}')
    print(f'Total revenue: {total:,.2f}')
    for k, v in sorted(by_channel.items()):
        print(f'  {k}: {v:,.2f}')
    for k, v in sorted(by_year.items()):
        print(f'  {k}: {v:,.2f}')
    for k, v in sorted(by_month.items()):
        print(f'  month {k}: {v:,.2f}')


if __name__ == '__main__':
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    main(*sys.argv[1:])
